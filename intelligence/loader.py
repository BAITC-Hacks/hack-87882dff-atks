"""Inspect the supplied System Electric workbooks without modifying them."""
from __future__ import annotations

import hashlib
import json
import math
import re
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd

DATA_DIR = Path(__file__).resolve().parent / "data"
MONTHS = {name: i for i, name in enumerate(
    ("янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"), 1)}
FILES = {
    "transactions": "Динамика продаж_Syseme Electric_2025-2026.xlsx",
    "monthly_sales": "Ежемесячные продажи в кол-м выражении SystemElectric 2024-2026.xlsx",
    "monthly_inventory": "Ежемесячные остатки SystemElectric 2024-2026.xlsx",
    "moq": "MOQ SystemElectric.xlsx",
    "goods_in_transit": "Товар в пути_SystemElectric на 22.09.2026.xlsx",
    "seasonality": "Сезонность SystemElectric 2024-2026.xlsx",
}
REQUIRED = {
    "transactions": {"Дата", "Номер", "Документ", "Код", "Номенклатура", "Ед.", "Склад", "Количество"},
    "monthly_sales": {"Номенклатура", "Номенклатура.Код", "Артикул", "Кратность"},
    "monthly_inventory": {"Номенклатура", "Номенклатура.Код", "Ед.изм"},
    "moq": {"Номенклатура", "Номенклатура.Код", "Артикул", "Кратность"},
    "goods_in_transit": {"Код 1с", "Артикул поставщика", "Наименование", "Категория 2026", "СЭ в пути 24.09"},
    "seasonality": {"год", *MONTHS},
}
RENAMES = {
    "Код": "sku", "Номенклатура.Код": "sku", "Код 1с": "sku",
    "Номенклатура": "product_name", "Наименование": "product_name",
    "Артикул": "article", "Артикул поставщика": "supplier_article",
    "Дата": "date", "Номер": "document_number", "Документ": "document",
    "Ед.": "unit", "Ед.изм": "unit", "Склад": "warehouse",
    "Количество": "quantity", "Кратность": "order_multiple",
    "Категория 2026": "category_2026", "СЭ в пути 24.09": "in_transit_quantity",
}


def clean(value):
    """Blank cells stay missing; literal strings such as NA are not null tokens."""
    if value is None or pd.isna(value):
        return pd.NA
    if isinstance(value, str):
        return value.strip() or pd.NA
    return value


def normalize_sku(value):
    value = clean(value)
    if pd.isna(value):
        return pd.NA
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if not math.isfinite(value):
            return pd.NA
        return str(int(value)) if value == int(value) else str(value)
    # Never strip leading zeroes, suffixes, or punctuation from text codes.
    return str(value)


def month_date(label):
    if isinstance(label, datetime):
        return pd.Timestamp(label.year, label.month, 1)
    match = re.fullmatch(r"([а-яё]+)\.?\s+(20\d{2})(?:\s+г\.)?", str(label).lower().strip())
    if match and match[1][:3] in MONTHS:
        return pd.Timestamp(int(match[2]), MONTHS[match[1][:3]], 1)
    return None


def _json_value(value):
    if value is pd.NA or value is pd.NaT:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if hasattr(value, "item"):
        return value.item()
    return str(value)


def _write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=_json_value,
                               allow_nan=False) + "\n", encoding="utf-8")


def inspect_workbook(path):
    """Report every sheet, including auxiliary sheets and uncached formulas."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    result = {"file": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "sheets": []}
    try:
        for sheet in book:
            headers, errors, formula_count = [], [], 0
            for row_number, row in enumerate(sheet.iter_rows(), 1):
                if row_number <= 10:
                    headers.append([c.value for c in row])
                for cell in row:
                    formula_count += cell.data_type == "f"
                    if cell.data_type == "e":
                        errors.append({"cell": cell.coordinate, "value": cell.value})
            result["sheets"].append({"sheet": sheet.title, "rows": sheet.max_row,
                                     "columns": sheet.max_column, "first_10_rows": headers,
                                     "formula_count": formula_count, "excel_errors": errors})
    finally:
        book.close()
    return result


def read_table(path, kind, audit):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    candidates = []
    try:
        for sheet in book:
            for row_number, row in enumerate(sheet.iter_rows(max_row=min(10, sheet.max_row)), 1):
                names = [str(c.value).strip() if c.value is not None else "" for c in row]
                if REQUIRED[kind].issubset(names):
                    candidates.append((sheet.title, row_number, names))
        if len(candidates) != 1:
            raise ValueError(f"{path.name}: expected one {kind} table, found {len(candidates)}")
        title, header_row, names = candidates[0]
        named = [n for n in names if n]
        if len(named) != len(set(named)):
            raise ValueError(f"{path.name}/{title}: duplicate column labels")
        # Positional labels preserve unnamed source columns; no field meaning is inferred.
        names = [n or f"source_column_{i}" for i, n in enumerate(names, 1)]
        rows = []
        for row_number, cells in enumerate(book[title].iter_rows(min_row=header_row + 1), header_row + 1):
            values = [clean(c.value) for c in cells]
            # Respect simple Excel zero-padding when a numeric SKU uses it.
            for i, name in enumerate(names):
                c = cells[i]
                if RENAMES.get(name) == "sku" and isinstance(c.value, (int, float)):
                    if re.fullmatch(r"0+", c.number_format) and float(c.value).is_integer():
                        values[i] = str(int(c.value)).zfill(len(c.number_format))
            rows.append([*values, row_number])
        frame = pd.DataFrame(rows, columns=[*names, "source_row"])
        blank = frame[names].isna().all(axis=1)
        total = pd.Series(False, index=frame.index)
        for col in ("Номенклатура", "Наименование", "Дата"):
            if col in frame:
                total |= frame[col].eq("Итого").fillna(False)
        secondary = pd.Series(False, index=frame.index)
        if kind == "monthly_sales":
            months = [c for c in names if month_date(c) is not None]
            secondary = frame[list(REQUIRED[kind])].isna().all(axis=1) & frame[months].eq("Количество").all(axis=1)
        exclude = blank | total | secondary
        audit.update({"file": path.name, "sheet": title, "header_row": header_row,
                      "source_columns": names, "excluded_blank_rows": int(blank.sum()),
                      "excluded_total_rows": frame.loc[total, "source_row"].tolist(),
                      "excluded_secondary_headers": frame.loc[secondary, "source_row"].tolist(),
                      "source_missing_values": frame.loc[~exclude, names].isna().sum().to_dict()})
        frame = frame.loc[~exclude].copy().rename(columns=RENAMES)
        if "sku" in frame:
            frame["sku"] = frame.sku.map(normalize_sku).astype("string")
        return frame
    finally:
        book.close()


def numeric(series, audit, column):
    cleaned = series.map(lambda v: v.replace("\u00a0", "").replace(" ", "").replace(",", ".")
                         if isinstance(v, str) else v)
    result = pd.to_numeric(cleaned, errors="coerce").astype("Float64")
    result = result.mask(result.isin([float("inf"), -float("inf")]))
    invalid = series.notna() & result.isna()
    audit.setdefault("invalid_numeric_values", {})[column] = {
        "count": int(invalid.sum()), "examples": series[invalid].head(10).tolist()}
    return result


def normalize_transactions(frame, audit):
    frame = frame.copy()
    frame["date_raw"] = frame.date
    def parse(value):
        if pd.isna(value):
            return pd.NaT
        if isinstance(value, (int, float)):
            return pd.to_datetime(value, unit="D", origin="1899-12-30", errors="coerce")
        return pd.to_datetime(value, dayfirst=True, errors="coerce")
    frame["date"] = pd.to_datetime(frame.date.map(parse))
    audit["invalid_dates"] = int((frame.date_raw.notna() & frame.date.isna()).sum())
    frame["quantity_raw"] = frame.quantity
    frame["quantity"] = numeric(frame.quantity, audit, "quantity")
    # Positive source movements are not outgoing sales. Preserve them for review.
    frame["demand"] = (-frame.quantity).where(frame.quantity <= 0)
    frame["movement_type"] = pd.Series("unknown", index=frame.index, dtype="string")
    frame.loc[frame.quantity < 0, "movement_type"] = "outgoing"
    frame.loc[frame.quantity > 0, "movement_type"] = "positive_movement"
    frame.loc[frame.quantity == 0, "movement_type"] = "zero"
    audit["movement_counts"] = frame.movement_type.value_counts().to_dict()
    return frame


def normalize_monthly(frame, value_name, audit):
    mapping = {c: month_date(c) for c in frame.columns if month_date(c) is not None}
    expected = set(pd.date_range("2024-01-01", "2026-09-01", freq="MS"))
    if set(mapping.values()) != expected or len(mapping) != len(expected):
        raise ValueError("Expected exactly one monthly column for each month Jan 2024–Sep 2026")
    result = frame.melt(id_vars=["sku", "source_row"], value_vars=list(mapping),
                        var_name="date", value_name=value_name)
    result["date"] = result.date.map(mapping)
    result[value_name] = numeric(result[value_name], audit, value_name)
    return result


def normalize_seasonality(frame, audit):
    # Only the top year/month table is the aggregate source. Other blocks are derived reports.
    years = pd.to_numeric(frame["год"], errors="coerce")
    mask = years.isin([2024, 2025, 2026])
    top = frame.loc[mask, ["год", *MONTHS, "source_row"]].copy()
    if len(top) != 3 or years[mask].nunique() != 3:
        raise ValueError("Expected one aggregate seasonality row per year 2024–2026")
    audit["excluded_auxiliary_rows"] = frame.loc[~mask, "source_row"].tolist()
    result = top.melt(id_vars=["год", "source_row"], value_vars=list(MONTHS),
                      var_name="month", value_name="value")
    result["date"] = [pd.Timestamp(int(y), MONTHS[m], 1) for y, m in zip(result["год"], result.month)]
    result["value"] = numeric(result.value, audit, "value")
    return result[["date", "value", "source_row"]]


def quality(frame, keys):
    content = frame.drop(columns=["source_row"], errors="ignore")
    result = {
        "row_count": len(frame), "sku_count": int(frame.sku.nunique()) if "sku" in frame else None,
        "missing_values": frame.isna().sum().to_dict(),
        "duplicate_count": int(content.duplicated().sum()),
        "duplicate_key_count": int(frame.duplicated(keys).sum()),
        "duplicate_key_columns": keys,
    }
    if "date" in frame:
        result["date_range"] = [None if pd.isna(v) else v.isoformat() for v in (frame.date.min(), frame.date.max())]
    for col in ("demand", "stock", "order_multiple", "in_transit_quantity", "value"):
        if col in frame:
            result[f"negative_{col}_count"] = int((frame[col] < 0).sum())
    if "order_multiple" in frame:
        result["nonpositive_order_multiple_count"] = int((frame.order_multiple <= 0).sum())
    return result


def run_pipeline(raw_dir=DATA_DIR / "raw", output_dir=DATA_DIR / "processed"):
    raw_dir, output_dir = Path(raw_dir).resolve(), Path(output_dir).resolve()
    if output_dir == raw_dir or raw_dir in output_dir.parents:
        raise ValueError("Output directory must not be inside raw data")
    paths = {kind: raw_dir / name for kind, name in FILES.items()}
    missing = [p.name for p in paths.values() if not p.is_file()]
    if missing:
        raise FileNotFoundError(f"Missing required workbooks: {missing}")
    output_dir.mkdir(parents=True, exist_ok=True)
    schema = {kind: inspect_workbook(path) for kind, path in paths.items()}
    _write_json(output_dir / "schema_report.json", schema)
    frames, audits, reports = {}, {}, {}
    for kind, path in paths.items():
        audit = audits[kind] = {}
        frame = read_table(path, kind, audit)
        if kind == "transactions":
            frame = normalize_transactions(frame, audit)
            keys = ["date", "document_number", "sku", "warehouse"]
        elif kind in ("monthly_sales", "monthly_inventory"):
            frame = normalize_monthly(frame, "demand" if kind == "monthly_sales" else "stock", audit)
            keys = ["sku", "date"]
        elif kind == "moq":
            frame["order_multiple"] = numeric(frame.order_multiple, audit, "order_multiple")
            keys = ["sku"]
        elif kind == "goods_in_transit":
            frame["in_transit_quantity"] = numeric(frame.in_transit_quantity, audit, "in_transit_quantity")
            keys = ["sku"]
        else:
            frame = normalize_seasonality(frame, audit)
            keys = ["date"]
        frames[kind] = frame
        reports[kind] = quality(frame, keys)
        reports[kind]["source"] = audit
    sku_sets = {k: set(f.sku.dropna()) for k, f in frames.items() if "sku" in f}
    tx = frames["transactions"]
    report = {
        "transaction_count": len(tx), "sku_count": len(set.union(*sku_sets.values())),
        "transaction_sku_count": len(sku_sets["transactions"]),
        "date_range": reports["transactions"]["date_range"],
        "warehouse_count": int(tx.warehouse.nunique()), "datasets": reports,
        "sku_coverage": {k: {"absent_from_monthly_sales": sorted(s - sku_sets["monthly_sales"]),
                              "monthly_sales_skus_absent_here": sorted(sku_sets["monthly_sales"] - s)}
                         for k, s in sku_sets.items()},
        "policies": ["No missing values filled; no duplicate records removed.",
                     "Transaction demand is -quantity for outgoing/zero movements; positive movements retain null demand.",
                     "Monthly sales are signed net values; negatives are retained and counted.",
                     "Кратность is an order multiple, not an independently verified minimum order quantity.",
                     "Seasonality values have unspecified units; no SKU or demand units inferred.",
                     "Excel formula cached values are used; formulas are not recalculated.",
                     "CSV nulls are empty fields; reload SKU columns explicitly as strings."],
    }
    # Validate the whole batch before replacing normalized data files.
    for kind, frame in frames.items():
        frame.to_csv(output_dir / f"{kind}.csv", index=False, encoding="utf-8", na_rep="")
    _write_json(output_dir / "quality_report.json", report)
    return report
